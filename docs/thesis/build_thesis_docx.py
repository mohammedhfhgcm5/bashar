# -*- coding: utf-8 -*-
"""Generate graduation thesis DOCX — chapters on new pages, abstracts, indexes, screenshots."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "اطروحة_مشروع_التخرج_منصة_توثيق_الأضرار.docx"
SMOKE = ROOT / "tmp" / "ui-smoke"

FIGURES = []  # (num, caption)
TABLES = []  # (num, caption)
_BOOKMARK_ID = 1


def set_rtl_run(run):
    rPr = run._element.get_or_add_rPr()
    if rPr.find(qn("w:rtl")) is None:
        rPr.append(OxmlElement("w:rtl"))


def set_rtl_para(p):
    pPr = p._element.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def style_run(run, size=14, bold=False, italic=False, english=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if english:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    else:
        run.font.name = "Traditional Arabic"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        set_rtl_run(run)


def _next_bookmark_id():
    global _BOOKMARK_ID
    i = _BOOKMARK_ID
    _BOOKMARK_ID += 1
    return i


def add_bookmark(paragraph, name: str):
    """Attach a Word bookmark to a paragraph so PAGEREF can resolve its page."""
    bid = str(_next_bookmark_id())
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _append_field(paragraph, instr: str, placeholder="—"):
    """Insert a Word field (PAGE / PAGEREF) into a paragraph."""
    run = paragraph.add_run()
    r = run._r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    r.append(begin)

    instr_run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    instr_run._r.append(it)

    sep_run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)

    result = paragraph.add_run(placeholder)
    style_run(result, size=12)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_page_field(paragraph):
    _append_field(paragraph, "PAGE", "1")


def add_page_ref_field(paragraph, bookmark: str):
    _append_field(paragraph, f"PAGEREF {bookmark} \\h", "—")


def setup_footer_page_numbers(doc: Document):
    """Centered page number in every section footer."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("— ")
        style_run(run, size=12)
        add_page_field(p)
        run2 = p.add_run(" —")
        style_run(run2, size=12)


def _shade_cell(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_col_width(cell, width_cm: float):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(int(width_cm * 567)))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def _write_cell_text(cell, text: str, *, bold=False, size=12, center=False, indent_cm=0.0, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    set_rtl_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    if indent_cm:
        pf.right_indent = Cm(indent_cm)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold)
    if color:
        run.font.color.rgb = color
    return p


def add_index_table(doc, rows, headers=("م", "العنوان", "الصفحة"), col_widths=(1.2, 13.5, 2.0)):
    """
    Pretty 3-column index table.
    rows: list of (serial, title, bookmark, is_section)
      is_section=True → bold + green tint (chapter/main entry)
    """
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    table.autofit = False

    header_fill = "1A3A2A"
    section_fill = "E8F0EA"
    alt_fill = "F7F9F7"
    white = "FFFFFF"

    # header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        _set_col_width(cell, col_widths[i])
        _shade_cell(cell, header_fill)
        _set_cell_margins(cell)
        _write_cell_text(
            cell,
            h,
            bold=True,
            size=13,
            center=True,
            color=RGBColor(0xFF, 0xFF, 0xFF),
        )

    for r_i, (serial, title, bookmark, is_section) in enumerate(rows):
        row = table.rows[r_i + 1]
        fill = section_fill if is_section else (alt_fill if r_i % 2 else white)

        # serial
        c0 = row.cells[0]
        _set_col_width(c0, col_widths[0])
        _shade_cell(c0, fill)
        _set_cell_margins(c0)
        _write_cell_text(c0, str(serial), bold=is_section, size=11, center=True)

        # title
        c1 = row.cells[1]
        _set_col_width(c1, col_widths[1])
        _shade_cell(c1, fill)
        _set_cell_margins(c1)
        indent = 0.0 if is_section else 0.35
        _write_cell_text(c1, title, bold=is_section, size=12 if is_section else 11, indent_cm=indent)

        # page (PAGEREF)
        c2 = row.cells[2]
        _set_col_width(c2, col_widths[2])
        _shade_cell(c2, fill)
        _set_cell_margins(c2)
        c2.text = ""
        p = c2.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_page_ref_field(p, bookmark)

    doc.add_paragraph()
    return table


def add_para(doc, text, size=14, bold=False, space_after=8, english=False, justify=True):
    p = doc.add_paragraph()
    set_rtl_para(p) if not english else None
    if english:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    elif justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, english=english)
    return p


def add_heading_rtl(doc, text, level=1, bookmark=None):
    h = doc.add_heading(level=level)
    set_rtl_para(h)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = h.add_run(text)
    sizes = {1: 18, 2: 16, 3: 14}
    style_run(run, size=sizes.get(level, 14), bold=True)
    if level == 1:
        run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x2A)
    if bookmark:
        add_bookmark(h, bookmark)
    return h


def add_list(doc, items, size=14):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_rtl_para(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(item)
        style_run(run, size=size)


def add_caption(doc, text, bookmark=None):
    p = doc.add_paragraph()
    set_rtl_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(text)
    style_run(run, size=12, bold=True)
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def add_table(doc, headers, rows, caption, table_id):
    TABLES.append((table_id, caption))
    bm = "bm_tbl_" + table_id.replace(".", "_")
    add_caption(doc, f"جدول {table_id}: {caption}", bookmark=bm)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        set_rtl_para(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        style_run(run, size=11, bold=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "DCE8E0")
        shading.set(qn("w:val"), "clear")
        cell._tc.get_or_add_tcPr().append(shading)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            set_rtl_para(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(val)
            style_run(run, size=11)
    doc.add_paragraph()
    return table


def add_figure(doc, path: Path, fig_id: str, caption: str, width_cm=14.5):
    FIGURES.append((fig_id, caption))
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width_cm))
    else:
        add_para(doc, f"[الصورة غير موجودة: {path.name}]", bold=True)
    bm = "bm_fig_" + fig_id.replace(".", "_")
    add_caption(doc, f"شكل {fig_id}: {caption}", bookmark=bm)


def page_break(doc):
    doc.add_page_break()


def chapter_title(doc, text, bookmark=None):
    page_break(doc)
    add_heading_rtl(doc, text, 1, bookmark=bookmark)


def update_fields_in_word(path: Path) -> bool:
    """Ask Microsoft Word to resolve PAGE / PAGEREF fields to real numbers."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        print("pywin32 not installed — fields will show after Update in Word")
        return False
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(path.resolve()))
        doc.Fields.Update()
        # some Word builds need a second pass for PAGEREF after PAGE
        doc.Repaginate()
        for story in doc.StoryRanges:
            story.Fields.Update()
        doc.Fields.Update()
        doc.Save()
        doc.Close(False)
        print("Word updated page fields in", path.name)
        return True
    except Exception as exc:
        print("Word field update skipped:", exc)
        return False
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def build():
    global _BOOKMARK_ID
    FIGURES.clear()
    TABLES.clear()
    _BOOKMARK_ID = 1

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.right_margin = Cm(2.8)
        section.left_margin = Cm(2.5)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    setup_footer_page_numbers(doc)

    # ========== TITLE ==========
    for _ in range(3):
        doc.add_paragraph()
    add_para(doc, "الجمهورية العربية السورية", size=16, bold=True, justify=False)
    add_para(doc, "مشروع تخرج", size=15, bold=True, justify=False)
    doc.add_paragraph()
    add_para(
        doc,
        "منصة توثيق الأضرار العقارية وتقييمها الأولي باستخدام الرؤية الحاسوبية",
        size=18,
        bold=True,
        justify=False,
    )
    add_para(doc, "(سجل الأضرار — Damage Log)", size=14, justify=False)
    doc.add_paragraph()
    add_para(doc, "إعداد الطالبين: [اسم الطالب الأول] — [اسم الطالب الثاني]", size=14, justify=False)
    add_para(doc, "بإشراف: [اسم المشرف]", size=14, justify=False)
    add_para(doc, "العام الدراسي 2025–2026", size=14, justify=False)

    # ========== ARABIC ABSTRACT ==========
    page_break(doc)
    add_heading_rtl(doc, "الملخص", 1, bookmark="bm_abstract")
    add_para(
        doc,
        "نقدّم في هذا المشروع منصة ويب متكاملة لتوثيق الأضرار العقارية بعد الكوارث أو الحوادث الموضعية. "
        "صمّمنا المنصة لتكون قناة واحدة تربط المواطن والمقيم بالمشرف ومنظمات الإغاثة: يُنشأ البلاغ بموقع جغرافي "
        "ووصف وصور، ثم يقدّم نموذج كشف أجسام من عائلة YOLOv8 تقييماً أولياً لمستوى الضرر، ويبقى قرار الاعتماد "
        "أو الرفض بيد المشرف البشري. بعد الاعتماد تستطيع المنظمة تسجيل مساعدات مخططة أو قيد التنفيذ أو مسلّمة "
        "وربطها بالبلاغ.",
    )
    add_para(
        doc,
        "بنينا النظام بهيكلية أحادية المستودع: واجهة React مع Vite وTailwind وLeaflet، وواجهة برمجية NestJS، "
        "وحزمة مشتركة للعقود وتجميع الكشف، وقاعدة PostgreSQL عبر Prisma. يعمل النموذج بصيغة ONNX داخل الخدمة "
        "دون الاعتماد على Python في التشغيل اليومي. اعتمدنا ستة أصناف عيوب من مجموعة PECCD، وتجميع كشف الصور "
        "المتعددة بقاعدة أشد عيب موثوق مع عتبة ثقة 0.35.",
    )
    add_para(
        doc,
        "الكلمات المفتاحية: توثيق الأضرار، YOLOv8، ONNX Runtime، نظم معلومات جغرافية، إدارة المساعدات، "
        "صلاحيات الأدوار، React، NestJS.",
        bold=True,
    )

    # ========== NUMBERED TOC (table + PAGEREF page numbers) ==========
    page_break(doc)
    add_heading_rtl(doc, "فهرس المحتويات", 1, bookmark="bm_toc")
    add_para(doc, "يعرض الجدول التالي أقسام الأطروحة مع رقم الصفحة المقابل لكل بند.", size=12, justify=False)
    # (serial, title, bookmark, is_main_section)
    toc_rows = [
        ("1", "الملخص", "bm_abstract", True),
        ("2", "فهرس المحتويات", "bm_toc", True),
        ("3", "فهرس الأشكال", "bm_fig_index", True),
        ("4", "فهرس الجداول", "bm_tbl_index", True),
        ("5", "الفصل الأول — المقدمة والإطار العام", "bm_ch1", True),
        ("1.1", "المقدمة", "bm_s1_1", False),
        ("1.2", "الهدف من المشروع", "bm_s1_2", False),
        ("1.3", "الأدوات المستخدمة", "bm_s1_3", False),
        ("1.4", "مزايا المشروع", "bm_s1_4", False),
        ("1.5", "مخطط سير المشروع", "bm_s1_5", False),
        ("6", "الفصل الثاني — الأدوات والتقنيات البرمجية", "bm_ch2", True),
        ("2.1", "بيئة التطوير", "bm_s2_1", False),
        ("2.2", "تقنيات الواجهة الأمامية (Frontend)", "bm_s2_2", False),
        ("2.3", "تقنيات الخادم (Backend)", "bm_s2_3", False),
        ("2.4", "تقنيات الرؤية الحاسوبية", "bm_s2_4", False),
        ("2.5", "متصفحات الويب وتقنيات العرض", "bm_s2_5", False),
        ("7", "الفصل الثالث — متطلبات النظام وتوصيفه", "bm_ch3", True),
        ("3.1", "المتطلبات الوظيفية", "bm_s3_1", False),
        ("3.2", "المتطلبات غير الوظيفية", "bm_s3_2", False),
        ("3.3", "توصيف النظام", "bm_s3_3", False),
        ("3.4", "النتائج المحققة", "bm_s3_4", False),
        ("8", "الفصل الرابع — القسم العملي — واجهات الموقع", "bm_ch4", True),
        ("4.1", "نظرة عامة على الواجهات", "bm_s4_1", False),
        ("4.2", "واجهات دور المواطن", "bm_s4_2", False),
        ("4.3", "واجهات دور المشرف", "bm_s4_3", False),
        ("4.4", "واجهات دور المنظمة", "bm_s4_4", False),
        ("4.5", "مسار الاستخدام المتكامل", "bm_s4_5", False),
        ("9", "الفصل الخامس — الآفاق المستقبلية والصعوبات", "bm_ch5", True),
        ("5.1", "الآفاق المستقبلية", "bm_s5_1", False),
        ("5.2", "الصعوبات التي واجهناها", "bm_s5_2", False),
        ("10", "المراجع والمصادر", "bm_refs", True),
        ("11", "Abstract (English)", "bm_abstract_en", True),
    ]
    add_index_table(doc, toc_rows)

    fig_index = [
        ("4.1", "صفحة قائمة البلاغات من حساب المواطن"),
        ("4.2", "تفاصيل البلاغ مع اقتراح النموذج وصناديق الكشف على الصور"),
        ("4.3", "محرر المسودة مع الخريطة التفاعلية وإرسال البلاغ"),
        ("4.4", "صفحة الخريطة لدى المواطن مع تجميع العلامات"),
        ("4.5", "صفحة المساعدات من حساب المشرف"),
        ("4.6", "طابور مراجعة البلاغات لدى المشرف"),
        ("4.7", "شاشة مراجعة المشرف (اعتماد / رفض / تعديل)"),
        ("4.8", "تسجيل توزيع مساعدة من حساب المنظمة"),
        ("4.9", "لوحة الإحصاءات لدى منسق المنظمة"),
    ]
    tbl_index = [
        ("1.1", "الأدوات والتقنيات المستخدمة في المشروع"),
        ("1.2", "مراحل سير تطوير المشروع"),
        ("3.1", "الأدوار والصلاحيات الرئيسية"),
        ("3.2", "بنية جداول قاعدة البيانات"),
        ("3.3", "وحدات الـ API ونقاط النهاية"),
        ("3.4", "النتائج الكمية المحققة"),
        ("4.1", "ربط لقطات الواجهة بالأدوار والمسارات"),
    ]

    page_break(doc)
    add_heading_rtl(doc, "فهرس الأشكال", 1, bookmark="bm_fig_index")
    add_para(doc, "قائمة الأشكال الواردة في الفصل الرابع مع أرقام صفحاتها.", size=12, justify=False)
    fig_rows = [
        (f"شكل {num}", cap, "bm_fig_" + num.replace(".", "_"), True)
        for num, cap in fig_index
    ]
    add_index_table(doc, fig_rows, headers=("الرقم", "عنوان الشكل", "الصفحة"), col_widths=(2.2, 12.5, 2.0))

    page_break(doc)
    add_heading_rtl(doc, "فهرس الجداول", 1, bookmark="bm_tbl_index")
    add_para(doc, "قائمة الجداول الواردة في الأطروحة مع أرقام صفحاتها.", size=12, justify=False)
    table_rows = [
        (f"جدول {num}", cap, "bm_tbl_" + num.replace(".", "_"), True)
        for num, cap in tbl_index
    ]
    add_index_table(doc, table_rows, headers=("الرقم", "عنوان الجدول", "الصفحة"), col_widths=(2.2, 12.5, 2.0))

    # ===================== CHAPTER 1 =====================
    chapter_title(doc, "الفصل الأول   —   المقدمة والإطار العام", bookmark="bm_ch1")

    add_heading_rtl(doc, "1.1  المقدمة", 2, bookmark="bm_s1_1")
    add_para(
        doc,
        "تُعدّ مشاريع التخرج ركيزةً أساسية في المسيرة الأكاديمية للطالب الجامعي، إذ تُتيح له فرصة توظيف ما اكتسبه "
        "من معارف ومهارات نظرية في بناء حلول تقنية حقيقية. وفي ظل التطور المتسارع في مجال تقنية المعلومات والرؤية "
        "الحاسوبية، باتت الحاجة ماسّةً إلى منصات رقمية متكاملة تُيسّر توثيق الأضرار العقارية وتنظّم مسارات المراجعة "
        "وتوزيع المساعدات.",
    )
    add_para(
        doc,
        "نقدّم في هذا المشروع منصةً إلكترونيةً شاملةً أسميناها عملياً «سجل الأضرار»، تستهدف تسهيل رفع البلاغات مع "
        "الموقع والصور، وتقديم تقييم أولي بمساعدة نموذج كشف أجسام، ومتابعة مراجعة المشرف، ثم تمكين منظمات الإغاثة "
        "من تسجيل المساعدات على البلاغات المعتمدة. عملنا على المشروع معاً كفريق ثنائي: صمّمنا العقود المشتركة، وبنينا "
        "الواجهة والواجهة البرمجية، وربطنا نموذج الرؤية بمسار البلاغ دون أن يستبدل القرار البشري.",
    )

    add_heading_rtl(doc, "1.2  الهدف من المشروع", 2, bookmark="bm_s1_2")
    add_para(doc, "نهدف من هذا المشروع إلى تحقيق جملة من الغايات الأكاديمية والتقنية:")
    add_list(
        doc,
        [
            "تطوير منصة ويب متكاملة تُنظّم دورة حياة بلاغ الضرر من المسودة حتى الاعتماد أو الرفض.",
            "توفير بيئة موحّدة تربط المواطن/المقيم بالمشرف ومنظمة الإغاثة بصلاحيات واضحة.",
            "أتمتة التقييم الأولي للصور عبر نموذج YOLOv8 يعمل بصيغة ONNX داخل الخادم، مع إبقاء الاعتماد بيد المشرف.",
            "بناء خريطة تفاعلية وإحصاءات تساعد على رؤية التجمع الجغرافي ومستويات الضرر.",
            "تمكين تسجيل المساعدات (غذاء، مأوى، طب، نقد، إعادة إعمار…) وربطها بالبلاغات المعتمدة.",
            "اكتساب خبرة عملية مشتركة في بناء تطبيق ويب بمستوى إنتاجي باستخدام تقنيات حديثة.",
        ],
    )

    add_heading_rtl(doc, "1.3  الأدوات المستخدمة", 2, bookmark="bm_s1_3")
    add_para(
        doc,
        "اعتمدنا في هذا المشروع على مجموعة من التقنيات الحديثة في تطوير تطبيقات الويب والرؤية الحاسوبية، "
        "نُلخّصها في الجدول التالي ونفصّل شرح كل منها في الفصل الثاني:",
    )
    add_table(
        doc,
        ["الأداة / التقنية", "الغرض", "الإصدار / الملاحظات"],
        [
            ["Visual Studio Code / Cursor", "محرر الأكواد وبيئة التطوير", "أحدث إصدار"],
            ["Node.js + TypeScript", "تشغيل الخادم والواجهة بلغات مطبَّقة الأنواع", "Node 20+"],
            ["React JS / Vite", "بناء واجهة المستخدم التفاعلية (SPA)", "React 18 / Vite 6"],
            ["Tailwind CSS", "تنسيق وتصميم واجهة المستخدم", "3.x"],
            ["React Router / React Query", "التنقل وإدارة طلبات الـ API", "Router 7 / Query 5"],
            ["Leaflet / React-Leaflet", "الخريطة التفاعلية للبلاغات", "Leaflet 1.9"],
            ["NestJS", "إطار بناء الـ Backend", "11.x"],
            ["PostgreSQL + Docker", "قاعدة البيانات العلاقية", "عبر Compose"],
            ["Prisma ORM", "الـ Schema والتهجير والوصول Type-safe", "6.x"],
            ["JWT + Passport + bcryptjs", "الهوية والجلسات", "—"],
            ["Multer + Sharp", "رفع الصور ومعالجتها", "—"],
            ["YOLOv8 + ONNX Runtime", "تقييم الأضرار من الصور", "yolov8n / ORT 1.30"],
            ["Zod + packages/shared", "العقود المشتركة", "Zod 3.x"],
            ["Jest", "اختبارات الوحدة", "—"],
        ],
        "الأدوات والتقنيات المستخدمة في المشروع",
        "1.1",
    )

    add_heading_rtl(doc, "1.4  مزايا المشروع", 2, bookmark="bm_s1_4")
    add_para(doc, "تتمتع المنصة بعدد من المزايا التي تميّزها عن الحلول التقليدية المتفرقة:")
    add_list(
        doc,
        [
            "دورة حياة بلاغ صريحة (مسودة ← مُرسَل ← قيد المراجعة ← معتمد / مرفوض) مع رمز مرجع فريد.",
            "تقييم أولي متعدد الصور بنموذج YOLOv8 مع ملخص عربي وثقة وتغطية تقريبية للعيوب.",
            "فصل واضح بين الاقتراح الآلي واعتماد المشرف.",
            "خريطة تفاعلية تعرض البلاغات وفق صلاحية الدور.",
            "نظام أدوار مع سياسة وصول على الخادم لا على الواجهة وحدها.",
            "سجل مراجعة يحفظ الاعتماد والرفض والتعديل.",
            "توزيع مساعدات مرتبط بالبلاغ المعتمد.",
            "واجهة عربية متجاوبة مبنية بـ React / Vite / Tailwind CSS.",
        ],
    )

    add_heading_rtl(doc, "1.5  مخطط سير المشروع", 2, bookmark="bm_s1_5")
    add_para(doc, "مرّ عملنا بمراحل تطوير متسلسلة ومترابطة:")
    add_table(
        doc,
        ["المرحلة", "الوصف"],
        [
            [
                "المرحلة الأولى — الأساس",
                "إعداد البيئة (NestJS + Prisma + PostgreSQL)، المصادقة JWT، الأدوار، واجهة الدخول، وبنية المستودع الأحادي.",
            ],
            [
                "المرحلة الثانية — اللب الوظيفي",
                "دورة البلاغ والصور، سياسة الصلاحيات، الخريطة، مراجعة المشرف، المنظمات والمساعدات، والإحصاءات.",
            ],
            [
                "المرحلة الثالثة — الرؤية والتكميل",
                "تدريب/تصدير YOLOv8 على PECCD، تشغيل ONNX، تجميع الكشف، الاختبارات، ولقطات التحقق من الواجهة.",
            ],
        ],
        "مراحل سير تطوير المشروع",
        "1.2",
    )

    add_para(
        doc,
        "خلاصة الفصل: استعرضنا في هذا الفصل الإطار العام للمشروع: الدافع، الأهداف، المزايا، جدول الأدوات، "
        "ومراحل التنفيذ. يُشكّل هذا الفصل المرجع الأساسي لفهم السياق الذي انطلقنا منه قبل الدخول في تفصيل التقنيات.",
        bold=True,
    )

    # ===================== CHAPTER 2 — EXPANDED =====================
    chapter_title(doc, "الفصل الثاني   —   الأدوات والتقنيات البرمجية", bookmark="bm_ch2")
    add_para(
        doc,
        "نفصّل في هذا الفصل كل أداة وتقنية اعتمدناها: ما هي، لماذا اخترناها، وكيف وظّفناها داخل منصة سجل الأضرار. "
        "الهدف أن يفهم القارئ ليس أسماء المكتبات فحسب، بل دور كل واحدة في حل مشكلة المشروع.",
    )

    add_heading_rtl(doc, "2.1  بيئة التطوير", 2, bookmark="bm_s2_1")

    add_heading_rtl(doc, "2.1.1  Visual Studio Code / Cursor", 3)
    add_para(
        doc,
        "Visual Studio Code محرر أكواد مفتوح المصدر من مايكروسوفت، يدعم الامتدادات وإكمال الأكواد وتكامل Git "
        "وتنقيح البرامج بعدة لغات. استخدمناه مع بيئة Cursor لتسريع التنقل بين ملفات المستودع الأحادي ومراجعة الفروقات. "
        "في عملنا اليومي فتحنا طرفيات متوازية: خادم NestJS، واجهة Vite، وحاوية PostgreSQL، مع الاعتماد على تمييز "
        "TypeScript للأخطاء قبل التشغيل.",
    )
    add_para(
        doc,
        "لماذا اخترناها؟ لأنها خفيفة، واسعة الدعم لمشاريع Node/React، وتناسب العمل الثنائي عبر Git دون فرض بيئة ثقيلة. "
        "وما الذي أنجزناه بها؟ كتابة وحدات الـ API ومكوّنات الواجهة والحزمة المشتركة من مساحة عمل واحدة.",
    )

    add_heading_rtl(doc, "2.1.2  Docker Compose", 3)
    add_para(
        doc,
        "Docker يعزل الخدمات في حاويات، وCompose يشغّل عدة خدمات بملف واحد. اعتمدناه لتشغيل PostgreSQL محلياً "
        "بأمر docker compose up -d، فتجنّبنا اختلاف إصدارات القاعدة بين أجهزتنا ووحّدنا بيانات التطوير والبذر. "
        "هذا الاختيار يقلّل عبارة «يعمل على جهازي» ويجعل إعادة بناء البيئة بعد أي خلل أمراً متوقعاً.",
    )

    add_heading_rtl(doc, "2.1.3  Node.js و TypeScript و npm workspaces", 3)
    add_para(
        doc,
        "Node.js منصة تشغيل JavaScript على الخادم، وTypeScript يضيف أنواعاً ثابتة تقلّل الأخطاء وقت التطوير. "
        "نظمنا المستودع كـ npm workspaces يضم apps/web وapps/api وpackages/shared، بحيث نبني الحزمة المشتركة أولاً "
        "ثم نستوردها في الطرفين. بهذا ضمنّا أن تسميات مستويات الضرر وأصناف العيوب ودالة التجميع واحدة في الواجهة والخادم.",
    )

    add_heading_rtl(doc, "2.2  تقنيات الواجهة الأمامية (Frontend)", 2, bookmark="bm_s2_2")

    add_heading_rtl(doc, "2.2.1  React JS", 3)
    add_para(
        doc,
        "React مكتبة لبناء واجهات تفاعلية عبر مكوّنات قابلة لإعادة الاستخدام وآلية Virtual DOM. بنينا بها كل الصفحات: "
        "الخريطة، قائمة البلاغات، التفصيل، الإحصاءات، المساعدات، وإدارة المستخدمين والمنظمات. اعتمدنا المكوّنات الوظيفية "
        "والـ Hooks (مثل useState وuseEffect عبر مزود المصادقة) لتقسيم المسؤولية: الغلاف Shell للتنقل، والصفحات للمنطق.",
    )
    add_para(
        doc,
        "اخترنا React لأن نظام الأدوار لدينا يغيّر القوائم والمسارات ديناميكياً، ولأن مكتبات الخرائط والنماذج "
        "(Leaflet، React Hook Form، React Query) تتكامل معها بسهولة. كما أن المجتمع والتوثيق سهّلا عملنا المشترك.",
    )

    add_heading_rtl(doc, "2.2.2  Vite", 3)
    add_para(
        doc,
        "Vite أداة بناء حديثة تعتمد وحدات ES الأصلية في المتصفح أثناء التطوير، مع Hot Module Replacement سريع جداً "
        "مقارنةً بـ Create React App. واجهتنا تعمل على المنفذ 5173 وتتصل بالـ API عبر أصل مضبوط وCORS. "
        "في الإنتاج يبني Vite حزمة محسّنة نخدمها كملفات ثابتة. اخترناه لأن سرعة الحلقة التطويرية كانت حاسمة ونحن نعدّل "
        "الواجهة والخادم معاً.",
    )

    add_heading_rtl(doc, "2.2.3  Tailwind CSS", 3)
    add_para(
        doc,
        "Tailwind إطار CSS بأسلوب Utility-First: نركّب المظهر من فئات صغيرة في JSX بدل ملفات CSS طويلة متفرقة. "
        "صمّمنا به الشريط الجانبي الداكن، منطقة المحتوى الفاتحة، شارات الحالة (مسودة / قيد المراجعة / معتمد)، "
        "وأزرار الاعتماد والرفض. دعم الاتجاه RTL يناسب العربية، ونظام الألوان الموحّد حافظ على اتساق الصفحات بين الأدوار.",
    )

    add_heading_rtl(doc, "2.2.4  React Router", 3)
    add_para(
        doc,
        "React Router يدير مسارات تطبيق الصفحة الواحدة دون إعادة تحميل كاملة. نفّذنا بوابة جلسة Gate تعيد غير المسجّل "
        "إلى /login، وبوابة دور RoleGate تمنع المواطن من صفحات الإدارة وتعيده إلى /map. المسارات تشمل /map و/reports "
        "و/reports/new و/reports/:id و/stats و/aid و/admin/users و/admin/organizations. هذا الفصل بين «الظهور في القائمة» "
        "و«السماح في المسار» يقلّل وصولاً غير مقصود حتى قبل رفض الخادم.",
    )

    add_heading_rtl(doc, "2.2.5  TanStack React Query", 3)
    add_para(
        doc,
        "React Query تدير دورة حياة طلبات الـ API: التخزين المؤقت، إعادة الجلب بعد التركيز على النافذة، وحالات "
        "التحميل والخطأ. استخدمناها لقوائم البلاغات والتقييم والإحصاءات والمساعدات حتى لا نكتب useEffect يدوي لكل صفحة. "
        "بعد إرسال بلاغ أو اعتماد أو تسجيل مساعدة نُبطل الاستعلامات ذات الصلة فتظهر البيانات الجديدة دون ارتباك في الحالة.",
    )

    add_heading_rtl(doc, "2.2.6  React Hook Form و Zod", 3)
    add_para(
        doc,
        "React Hook Form تقلّل إعادة تصيير النماذج الكبيرة (إنشاء بلاغ، مراجعة، تسجيل مساعدة). Zod يصف مخطط التحقق "
        "بقواعد صريحة (بريد، كلمة مرور، تعدادات). ربطناهما عبر @hookform/resolvers، وأعدنا استخدام Zod في الحزمة "
        "المشتركة والخادم قدر الإمكان حتى لا تختلف رسالة الخطأ بين الطرفين. هذا يقلّل بلاغات ناقصة أو بحقول خارج التعدادات.",
    )

    add_heading_rtl(doc, "2.2.7  Leaflet و React-Leaflet", 3)
    add_para(
        doc,
        "Leaflet مكتبة خرائط خفيفة مفتوحة المصدر، وReact-Leaflet تلفّها بمكوّنات React. أضفنا تجميع العلامات "
        "(marker clustering) حتى لا تتراكب النقاط عند الزوم البعيد. بلاط الخريطة من OpenStreetMap. "
        "في صفحة الخريطة يعرض المواطن بلاغاته ضمن الإطار، ويختار موقع العقار بالنقر على الخريطة داخل محرر المسودة. "
        "لماذا Leaflet؟ لأنها كافية لعلامات وتصفية دون تكلفة خرائط تجارية، وتتكامل مع إحداثيات lat/lng المخزّنة في Property.",
    )

    add_heading_rtl(doc, "2.3  تقنيات الخادم (Backend)", 2, bookmark="bm_s2_3")

    add_heading_rtl(doc, "2.3.1  NestJS", 3)
    add_para(
        doc,
        "NestJS إطار Backend بـ TypeScript على Node.js، ببنية مستلهمة من Angular: Modules وControllers وServices "
        "وحقن تبعيات وتزيينات. نظّمنا العمل في وحدات: Auth، Users، Organizations، Reports، Vision، Aid، Stats، Storage. "
        "كل مسار حساس يمر بحارس JWT ثم حارس أدوار عند الحاجة. اخترنا NestJS لأن تعقيد RBAC ودورة البلاغ يحتاج تنظيماً "
        "أوضح مما يقدّمه Express العاري، ولأن الاختبار وحقن الخدمات أسهل في هذه البنية.",
    )

    add_heading_rtl(doc, "2.3.2  PostgreSQL", 3)
    add_para(
        doc,
        "PostgreSQL نظام قواعد بيانات علاقي موثوق يدعم قيوداً قوية وفهارس وأنواعاً متقدمة. خزّنا فيه المستخدمين، "
        "المنظمات، العقارات بإحداثياتها، البلاغات وحالاتها، الصور الوصفية، تقييم الرؤية، سجل المراجعة، والمساعدات. "
        "وضعنا فهارس على status وdamageLevel و(lat,lng) وreferenceCode لتسريع التصفية والخريطة. اخترناه لجدية القيود "
        "والعلاقات بدل قاعدة وثائقية قد تُضعف تكامل دورة الاعتماد والمساعدة.",
    )

    add_heading_rtl(doc, "2.3.3  Prisma ORM", 3)
    add_para(
        doc,
        "Prisma يوفّر مخططاً تعريفياً schema.prisma، وعميلاً Type-safe، وأداة Migrate للتهجير. عرّفنا التعدادات "
        "(Role، ReportStatus، DamageLevel، AidType، VisionStatus…) والعلاقات بين Report وImage وVisionAssessment. "
        "سهّل علينا البذر التجريبي والتهجير المشترك بين جهازينا. بدل كتابة SQL يدوي لكل عملية، نكتب استعلامات واضحة "
        "مع أمان أنواع يقلّل أخطاء الحقول وقت الترجمة.",
    )

    add_heading_rtl(doc, "2.3.4  JWT و Passport و bcryptjs", 3)
    add_para(
        doc,
        "المصادقة عندنا على ثلاث طبقات: bcryptjs لتجزئة كلمة المرور عند التسجيل والتحقق عند الدخول؛ JWT لرمز وصول "
        "قصير العمر يحمله العميل؛ Passport-JWT للتحقق من الرمز في الحراس. رموز التحديث تُحفظ كبصمة في RefreshSession "
        "مع تاريخ انتهاء، ويمكن إلغاؤها عند الخروج. ضبطنا CORS على أصل الواجهة مع credentials حتى تعمل الكوكيز/الجلسة "
        "بأمان في التطوير. الهدف: لا كلمات مرور نصية، ولا رمز تحديث قابل للسرقة من قاعدة البيانات كنص صريح.",
    )

    add_heading_rtl(doc, "2.3.5  Multer و Sharp", 3)
    add_para(
        doc,
        "Multer وسيط لرفع الملفات عبر multipart في نقاط نهاية صور البلاغ. Sharp مكتبة معالجة صور عالية الأداء: "
        "نستخدمها لإعادة التحجيم بأسلوب letterbox إلى 640×640 قبل الاستدلال، بما يحافظ على نسبة العرض إلى الارتفاع. "
        "كما فحصنا محتوى الملف (image sniff) حتى لا يُقبل مستند باسم صورة. بهذا نربط مسار التخزين المحلي بمفتاح storageKey "
        "بدون وضع الملف الثنائي داخل صف قاعدة البيانات.",
    )

    add_heading_rtl(doc, "2.3.6  الحزمة المشتركة packages/shared", 3)
    add_para(
        doc,
        "أنشأنا حزمة داخل المستودع تصدّر الأدوار، تسميات المستويات، أصناف عيوب PECCD الستة، ودالة aggregateDamage. "
        "الواجهة تعرض الملخص العربي نفسه الذي يحسبه الخادم عند التقييم. عتبة VISION_MIN_CONFIDENCE = 0.35 موحّدة. "
        "هذا القرار المعماري منع انحرافاً شائعاً في المشاريع الثنائية: نص مختلف أو قاعدة شدة مختلفة بين Frontend وBackend.",
    )

    add_heading_rtl(doc, "2.3.7  Jest", 3)
    add_para(
        doc,
        "Jest إطار اختبارات وحدة. كتبنا اختبارات لتجميع الضرر، وسياسة قراءة/إرسال البلاغ، ومسار التخزين وفحص الصور. "
        "هذه الاختبارات تحمي القواعد الحساسة من الانكسار عند أي تعديل لاحق، وهي جزء من جودة المشروع لا من التجميل.",
    )

    add_heading_rtl(doc, "2.4  تقنيات الرؤية الحاسوبية", 2, bookmark="bm_s2_4")

    add_heading_rtl(doc, "2.4.1  YOLOv8 (Ultralytics)", 3)
    add_para(
        doc,
        "YOLOv8 عائلة كاشفات أجسام تعالج الصورة في مرور أمامي واحد وتعيد صنفاً ومستطيلاً محيطاً وثقة. درّبنا النسخة "
        "الصغيرة YOLOv8n عبر Ultralytics على تقسيم PECCD (حوالي عشرات الحقب مع early stopping)، لأن هدفنا فرز ميداني "
        "على خادم بلا بطاقة رسوم مخصّصة للاستدلال لا بناء نموذج بحثي ضخم. الأصناف الستة: شق عميق، شق بسيط، شق متفرع، "
        "تقشر خرسانة، تفتت سطحي، وثقوب — ولكل صنف شدة مرتبطة في المنتج (من منخفض إلى حرج).",
    )

    add_heading_rtl(doc, "2.4.2  ONNX Runtime", 3)
    add_para(
        doc,
        "بعد التدريب نصدّر الأفضل إلى peccd-yolov8n.onnx. في الإنتاج يحمّل NestJS الملف عبر onnxruntime-node بمزود CPU. "
        "بهذا فصلنا بيئة التدريب Python عن التشغيل اليومي Node. إن غاب الملف تُحفظ حالة UNAVAILABLE؛ وإن فشل الاستدلال "
        "FAILED — دون إسقاط بقية الـ API. اسم النموذج المخزَّن مع التقييم: yolov8n-peccd.",
    )

    add_heading_rtl(doc, "2.4.3  مجموعة بيانات PECCD", 3)
    add_para(
        doc,
        "PECCD مجموعة بيانات عيوب خرسانة ظاهرة (Mayya وآخرون، Mendeley Data، 2025). اعتمدنا ترتيب أصنافها كما هو "
        "مثبّت في الحزمة المشتركة. المنصة لا تدّعي أن النموذج يغطي الحريق أو الانهيار الكامل أو العيوب غير المرئية؛ "
        "هي أداة فرز للعيوب الظاهرة ضمن هذه الأصناف فقط.",
    )

    add_heading_rtl(doc, "2.4.4  قاعدة تجميع الضرر", 3)
    add_para(
        doc,
        "بعد كشف كل صور البلاغ: نُسقط ما دون ثقة 0.35، نجمّع العدد وأعلى ثقة لكل صنف، ثم المستوى المقترح هو أعلى شدة "
        "ظهرت (لا المتوسط). الثقة المعروضة تُشتق من أعلى كشف داعم مع زيادة طفيفة عند اتفاق عدة صور (سقف 0.99). "
        "نسبة التغطية تقريبية على شبكة 64×64 باتحاد الخلايا. النتيجة تظهر في الواجهة كاقتراح يمكن تطبيقه يدوياً على "
        "حقل مستوى الضرر — دون كتابة فوق اختيار المبلّغ تلقائياً عند الإنشاء.",
    )

    add_heading_rtl(doc, "2.5  متصفحات الويب وتقنيات العرض", 2, bookmark="bm_s2_5")
    add_para(
        doc,
        "يعمل تطبيقنا كـ Single Page Application: المتصفح يحمّل HTML/CSS/JS مرة ثم يدير التنقل عبر React Router. "
        "الطلبات إلى /api تُنجز عبر دوال العميل مع إدارة React Query. المتصفحات المستهدفة: Chrome وEdge وFirefox وSafari "
        "الحديثة التي تدعم ES modules وFetch وCanvas/الخرائط.",
    )

    add_para(
        doc,
        "خلاصة الفصل: قدّمنا منظومة تقنية مترابطة — من بيئة التطوير حتى مسار الرؤية — مع تبرير كل اختيار بربطه "
        "بحاجة المنصة (أدوار، صور، خريطة، فرز آلي، مساعدات). التفاصيل العملية للواجهات في الفصل الرابع مع اللقطات المرقّمة.",
        bold=True,
    )

    # ===================== CHAPTER 3 =====================
    chapter_title(doc, "الفصل الثالث   —   متطلبات النظام وتوصيفه", bookmark="bm_ch3")

    add_heading_rtl(doc, "3.1  المتطلبات الوظيفية (Functional Requirements)", 2, bookmark="bm_s3_1")
    add_para(doc, "حددنا المتطلبات وفق الأدوار الأربعة كما يلي:")

    add_heading_rtl(doc, "3.1.1  متطلبات المواطن / المقيم", 3)
    add_list(
        doc,
        [
            "إنشاء حساب وتسجيل الدخول والخروج وتجديد الجلسة.",
            "إنشاء بلاغ مسودة (عنوان، إحداثيات، نوع عقار، وصف، مستوى ضرر).",
            "إضافة صور وحذفها ضمن السياسة.",
            "إرسال البلاغ ومتابعة حالته ورؤيته على الخريطة ضمن بلاغاته.",
            "طلب التقييم الآلي وتطبيق المستوى المقترح.",
            "تعديل المرفوض وإعادة الإرسال.",
        ],
    )

    add_heading_rtl(doc, "3.1.2  متطلبات المشرف", 3)
    add_list(
        doc,
        [
            "استعراض كل البلاغات وتصفيةها (حالة، مستوى، نوع، فترة، رمز).",
            "فتح بلاغ مُرسَل → قيد المراجعة.",
            "اعتماد أو رفض مع سبب أو تعديل مع سجل.",
            "تشغيل التقييم الآلي.",
            "إدارة المستخدمين والمنظمات.",
            "عرض الإحصاءات والمساعدات.",
        ],
    )

    add_heading_rtl(doc, "3.1.3  متطلبات المنظمة", 3)
    add_list(
        doc,
        [
            "استعراض البلاغات المعتمدة فقط.",
            "تسجيل مساعدة على بلاغ معتمد وتحديث حالتها.",
            "عرض الإحصاءات المتاحة للمنظمة.",
        ],
    )

    add_heading_rtl(doc, "3.2  المتطلبات غير الوظيفية", 2, bookmark="bm_s3_2")
    add_list(
        doc,
        [
            "الأداء: استجابة واجهة مقبولة، واستدلال رؤية لا يُسقط بقية الخدمة.",
            "الأمان: bcryptjs، JWT، بصمات التحديث، RBAC على الخادم.",
            "التوسعية: وحدات NestJS وحزمة shared.",
            "سهولة الاستخدام: واجهة عربية RTL ورسائل واضحة.",
            "الاعتمادية: حالات UNAVAILABLE/FAILED للرؤية.",
            "التوافق: المتصفحات الحديثة الرئيسية.",
        ],
    )

    add_heading_rtl(doc, "3.3  توصيف النظام", 2, bookmark="bm_s3_3")
    add_para(
        doc,
        "النظام Client–Server داخل مستودع واحد: React، NestJS، shared، PostgreSQL، وملف ONNX للرؤية.",
    )

    add_heading_rtl(doc, "3.3.1  الأدوار والصلاحيات", 3)
    add_table(
        doc,
        ["الدور", "الصلاحيات والوظائف الرئيسية"],
        [
            [
                "مواطن / مقيم",
                "إنشاء بلاغاته وتعديل المسودة/المرفوض، الصور، الإرسال، التقييم، الخريطة الخاصة به.",
            ],
            [
                "مشرف",
                "كل البلاغات، المراجعة، إدارة المستخدمين والمنظمات، الإحصاء، الاطلاع على المساعدات.",
            ],
            [
                "منظمة",
                "البلاغات المعتمدة فقط، إنشاء المساعدات وتحديثها، الإحصاءات.",
            ],
        ],
        "الأدوار والصلاحيات الرئيسية",
        "3.1",
    )

    add_heading_rtl(doc, "3.3.2  بنية قاعدة البيانات", 3)
    add_table(
        doc,
        ["اسم الجدول", "الحقول الرئيسية", "الوصف"],
        [
            ["User", "id, fullName, email, passwordHash, role, organizationId", "المستخدمون والأدوار"],
            ["RefreshSession", "userId, tokenHash, expiresAt", "جلسات التحديث"],
            ["Organization", "name, contactEmail, contactPhone, isActive", "جهات الإغاثة"],
            ["Property", "address, lat, lng, propertyType", "العقار والموقع"],
            ["Report", "referenceCode, status, damageLevel, description…", "البلاغ ودورة حياته"],
            ["Image", "reportId, storageKey, mimeType, sizeBytes", "صور البلاغ"],
            ["VisionAssessment", "status, suggestedLevel, confidence, findings", "التقييم الآلي"],
            ["ReviewLog", "action, before, after, note", "سجل المراجعة"],
            ["AidDistribution", "aidType, status, organizationId, reportId", "المساعدات"],
            ["ReportCounter", "year, last", "عداد رمز المرجع"],
        ],
        "بنية جداول قاعدة البيانات",
        "3.2",
    )

    add_heading_rtl(doc, "3.3.3  API (Backend Endpoints)", 3)
    add_table(
        doc,
        ["الوحدة", "نقاط النهاية"],
        [
            ["AuthModule", "POST register/login/refresh/logout — GET me"],
            ["UsersModule", "GET/POST /users — PATCH /users/:id"],
            ["OrganizationsModule", "GET/POST /organizations — PATCH /:id"],
            ["ReportsModule", "قائمة، إنشاء، تعديل، submit، assess، review، images، map"],
            ["AidModule", "GET/POST /aid — PATCH /aid/:id"],
            ["StatsModule", "GET /stats/overview"],
            ["Health", "GET /api/health"],
        ],
        "وحدات الـ API ونقاط النهاية",
        "3.3",
    )

    add_heading_rtl(doc, "3.4  النتائج المحققة", 2, bookmark="bm_s3_4")
    add_table(
        doc,
        ["المعيار", "النتيجة المحققة"],
        [
            ["وحدات Backend", "مصادقة، مستخدمون، منظمات، بلاغات، رؤية، مساعدة، إحصاء، تخزين"],
            ["جداول قاعدة البيانات", "10 كيانات بعلاقات Prisma"],
            ["صفحات الواجهة الموثّقة بلقطات", "9 لقطات تغطي المواطن والمشرف والمنظمة"],
            ["الأدوار", "مواطن، مقيم، مشرف، منظمة"],
            ["أصناف الرؤية", "6 أصناف PECCD مع تجميع متعدد الصور"],
            ["المصادقة", "JWT + bcryptjs + RefreshSession + RBAC"],
        ],
        "النتائج الكمية المحققة",
        "3.4",
    )

    add_para(
        doc,
        "خلاصة الفصل: حلّلنا المتطلبات الوظيفية وغير الوظيفية، ووصفنا الأدوار وقاعدة البيانات والـ API، "
        "ولخّصنا النتائج التي تعكس اكتمال المنصة في نطاق مشروع التخرج.",
        bold=True,
    )

    # ===================== CHAPTER 4 — SCREENSHOTS =====================
    chapter_title(doc, "الفصل الرابع   —   القسم العملي — واجهات الموقع", bookmark="bm_ch4")

    add_heading_rtl(doc, "4.1  نظرة عامة على الواجهات", 2, bookmark="bm_s4_1")
    add_para(
        doc,
        "صمّمنا الواجهات بأسلوب RTL عربي موحّد: شريط جانبي داكن للتنقل حسب الدور، ومنطقة محتوى فاتحة للعمل. "
        "اللقطات التالية أُخذت من تشغيل تحقق الواجهة (ui-smoke) بعد بذر بيانات تجريبية، وهي تمثّل الصفحات الفعلية "
        "للمنصة لا نماذج تصميم منفصلة. نعرض لكل شكل: رقمها، ماذا تظهر، ولمن، وكيف تُستخدم في مسار العمل.",
    )
    add_table(
        doc,
        ["الشكل", "الملف المصدري", "الدور", "المسار / الاستخدام"],
        [
            ["4.1", "01-citizen-reports.png", "مواطن", "/reports — متابعة بلاغاته"],
            ["4.2", "02-citizen-report-assessment.png", "مواطن", "/reports/:id — تقييم الصور"],
            ["4.3", "03-citizen-draft-editor.png", "مواطن", "/reports/:id — تعديل مسودة وإرسال"],
            ["4.4", "04-citizen-map.png", "مواطن", "/map — العرض الجغرافي"],
            ["4.5", "05-supervisor-aid.png", "مشرف", "/aid — الاطلاع على المساعدات"],
            ["4.6", "06-supervisor-reports.png", "مشرف", "/reports — طابور المراجعة"],
            ["4.7", "07-supervisor-review.png", "مشرف", "/reports/:id — اعتماد/رفض"],
            ["4.8", "08-org-aid.png", "منظمة", "/aid — تسجيل توزيع"],
            ["4.9", "09-org-stats.png", "منظمة", "/stats — لوحة الإحصاء"],
        ],
        "ربط لقطات الواجهة بالأدوار والمسارات",
        "4.1",
    )

    # ---- Citizen ----
    add_heading_rtl(doc, "4.2  واجهات دور المواطن", 2, bookmark="bm_s4_2")

    add_heading_rtl(doc, "4.2.1  قائمة البلاغات", 3)
    add_para(
        doc,
        "تُعدّ صفحة البلاغات نقطة المتابعة اليومية للمواطن. في الشكل 4.1 يظهر الشريط الجانبي بروابط الخريطة والبلاغات "
        "وبلاغ جديد، مع هوية «مواطن تجريبي». أعلى الصفحة شريط تصفية: فترة زمنية، رمز البلاغ، نوع العقار، مستوى الضرر، "
        "والحالة. الجدول يعرض الرمز (مثل RPT-2026-00008) والعنوان ومستوى الضرر بلون مميّز والحالة (مسودة / قيد المراجعة / "
        "معتمد) ومقدّم البلاغ. الاستخدام: يفتح المواطن بلاغاً لمتابعة التقييم أو الإرسال أو معرفة نتيجة المراجعة.",
    )
    add_figure(
        doc,
        SMOKE / "01-citizen-reports.png",
        "4.1",
        "صفحة قائمة البلاغات من حساب المواطن — تصفية ومتابعة الحالات",
    )

    add_heading_rtl(doc, "4.2.2  تفاصيل البلاغ والتقييم الآلي", 3)
    add_para(
        doc,
        "الشكل 4.2 هو أهم لقطة تقنية في مسار المواطن: تفاصيل بلاغ مسودة مع قسم «اقتراح النموذج». بعد الضغط على "
        "«تقييم الصور» يجمع الخادم صور البلاغ، يشغّل نموذج YOLOv8 عبر ONNX، ثم يعيد ملخصاً عربياً (مثلاً: تفتت سطحي، "
        "مستوى مقترح متوسط، ثقة حو 77٪، تغطية عيوب حو 55٪). على الصور تظهر صناديق الكشف مع تسمية الصنف ونسبة الثقة. "
        "الاستخدام: يراجع المواطن الاقتراح قبل الإرسال، ويمكنه لاحقاً تطبيق المستوى المقترح على حقل الضرر. التقييم لا "
        "يغيّر حالة الاعتماد؛ هو أداة مساعدة فقط.",
    )
    add_figure(
        doc,
        SMOKE / "02-citizen-report-assessment.png",
        "4.2",
        "تفاصيل البلاغ مع اقتراح النموذج وصناديق الكشف على الصور",
        width_cm=14.0,
    )

    add_heading_rtl(doc, "4.2.3  محرر المسودة والإرسال", 3)
    add_para(
        doc,
        "الشكل 4.3 يكمل صفحة التفصيل من جهة التعديل: خريطة Leaflet لتحديد موقع العقار بالنقر، حقول العنوان ونوع العقار "
        "والوصف ومستوى الضرر، زر «استخدام المستوى المقترح»، وإضافة صور، ثم «حفظ» أو «إرسال البلاغ». الاستخدام: يُكمّل "
        "المواطن المسودة حتى تصبح جاهزة، ثم يرسلها فتنتقل إلى حالة مُرسَل لتظهر في طابور المشرف.",
    )
    add_figure(
        doc,
        SMOKE / "03-citizen-draft-editor.png",
        "4.3",
        "محرر المسودة مع الخريطة التفاعلية وأزرار الحفظ والإرسال",
        width_cm=14.0,
    )

    add_heading_rtl(doc, "4.2.4  صفحة الخريطة", 3)
    add_para(
        doc,
        "الشكل 4.4 يعرض الصفحة الافتراضية بعد الدخول: خريطة OpenStreetMap عبر Leaflet، مع مفتاح ألوان للمستويات "
        "(منخفض / متوسط / مرتفع / حرج)، وتصفية بنفس أبعاد قائمة البلاغات. العلامة المجمّعة (مثلاً الرقم 5) تعني عدة "
        "بلاغات متقاربة جغرافياً. الاستخدام: يفهم المواطن انتشار بلاغاته مكانياً وينتقل منها إلى التفاصيل.",
    )
    add_figure(
        doc,
        SMOKE / "04-citizen-map.png",
        "4.4",
        "صفحة الخريطة لدى المواطن مع تجميع العلامات ومفتاح مستويات الضرر",
    )

    # ---- Supervisor ----
    add_heading_rtl(doc, "4.3  واجهات دور المشرف", 2, bookmark="bm_s4_3")

    add_heading_rtl(doc, "4.3.1  المساعدات من منظور المشرف", 3)
    add_para(
        doc,
        "في الشكل 4.5 يتّسع الشريط الجانبي ليشمل الإحصاء والمساعدات والمستخدمين والمنظمات. صفحة المساعدات تعرض توزيعات "
        "المنظمات على البلاغات المعتمدة (مثال: إيواء/خيمة على RPT-2026-00001 بحالة تم التسليم ومستوى حرج). "
        "الاستخدام: يراقب المشرف وصول المساعدة دون أن ينشئها بدلاً عن المنظمة في المسار الاعتيادي.",
    )
    add_figure(
        doc,
        SMOKE / "05-supervisor-aid.png",
        "4.5",
        "صفحة المساعدات من حساب المشرف — متابعة التوزيعات على البلاغات المعتمدة",
    )

    add_heading_rtl(doc, "4.3.2  طابور مراجعة البلاغات", 3)
    add_para(
        doc,
        "الشكل 4.6 يظهر قائمة أوسع للمشرف («طابور المراجعة والبحث») مع بلاغات من مقدّمين مختلفين وحالات متعددة. "
        "التصفية تتيح عزل ما هو قيد المراجعة أو المعتمد. الاستخدام: يختار المشرف بلاغاً مُرسَلاً أو قيد المراجعة لفتحه.",
    )
    add_figure(
        doc,
        SMOKE / "06-supervisor-reports.png",
        "4.6",
        "طابور مراجعة البلاغات لدى المشرف مع التصفية والشارات الملونة",
    )

    add_heading_rtl(doc, "4.3.3  شاشة المراجعة والقرار", 3)
    add_para(
        doc,
        "الشكل 4.7 هو قلب صلاحية المشرف: بلاغ في حالة «قيد المراجعة» مع وصف وصور وإمكانية تشغيل التقييم إن لم يُنفَّذ. "
        "قسم المراجعة يتيح تعديل الوصف ومستوى الضرر وإضافة ملاحظة، ثم ثلاثة قرارات: اعتماد، حفظ التعديل، أو رفض. "
        "الاستخدام: يغلق المشرف دورة التوثيق بقرار بشري صريح يُسجَّل في ReviewLog، وبعد الاعتماد يظهر البلاغ للمنظمة.",
    )
    add_figure(
        doc,
        SMOKE / "07-supervisor-review.png",
        "4.7",
        "شاشة مراجعة المشرف — اعتماد أو رفض أو حفظ التعديل",
    )

    # ---- Organization ----
    add_heading_rtl(doc, "4.4  واجهات دور المنظمة", 2, bookmark="bm_s4_4")

    add_heading_rtl(doc, "4.4.1  تسجيل توزيع المساعدة", 3)
    add_para(
        doc,
        "الشكل 4.8 يخص منسق المنظمة: القائمة تقتصر على الخريطة والبلاغات المعتمدة والإحصاء والمساعدات. أعلى الصفحة "
        "نموذج تسجيل توزيع (البحث برمز البلاغ، اختيار بلاغ معتمد، نوع المساعدة، موعد، ملاحظات). أسفلها بطاقة توزيع "
        "مع أزرار الحالة: مخططة / قيد التنفيذ / تم التسليم / ملغاة. الاستخدام: لا تُنشأ مساعدة إلا على بلاغ معتمد، "
        "وهذا يربط الإغاثة بقرار المشرف لا بمسودات المواطنين.",
    )
    add_figure(
        doc,
        SMOKE / "08-org-aid.png",
        "4.8",
        "تسجيل توزيع مساعدة من حساب المنظمة وتحديث حالة التسليم",
    )

    add_heading_rtl(doc, "4.4.2  لوحة الإحصاءات", 3)
    add_para(
        doc,
        "الشكل 4.9 يعرض بطاقات إحصائية: حسب مستوى الضرر، حسب الحالة، المساعدات، حسب نوع العقار، ونشاط آخر 30 يوماً. "
        "الاستخدام: يقرأ منسق المنظمة صورة سريعة للاحتياج المعتمد وما تم تسليمه دون الدخول في بلاغات غير معتمدة.",
    )
    add_figure(
        doc,
        SMOKE / "09-org-stats.png",
        "4.9",
        "لوحة الإحصاءات لدى منسق المنظمة — الضرر والحالة والمساعدات ونوع العقار",
    )

    add_heading_rtl(doc, "4.5  مسار الاستخدام المتكامل", 2, bookmark="bm_s4_5")
    add_para(
        doc,
        "بجمع اللقطات يتضح المسار الذي صمّمناه: المواطن ينشئ ويتابع على القائمة والخريطة (4.1 و4.4)، يقيّم الصور "
        "ويكمل المسودة ثم يرسل (4.2 و4.3)؛ المشرف يراجع من الطابور ويقرر (4.6 و4.7) ويراقب المساعدات (4.5)؛ المنظمة "
        "تسجّل التوزيع وتحدّث حالته وتقرأ الإحصاء (4.8 و4.9). هذا التكامل هو جوهر القسم العملي للمشروع.",
    )

    add_para(
        doc,
        "خلاصة الفصل: وثّقنا الواجهات الفعلية بأرقام أشكال وشرح وظيفة كل صفحة وربطها بالدور والمسار، مع إبراز "
        "التقييم الآلي في الشكل 4.2 وقرار المشرف في الشكل 4.7 بوصفهما نقطتي الارتكاز التقني والإداري.",
        bold=True,
    )

    # ===================== CHAPTER 5 =====================
    chapter_title(doc, "الفصل الخامس   —   الآفاق المستقبلية والصعوبات", bookmark="bm_ch5")

    add_heading_rtl(doc, "5.1  الآفاق المستقبلية", 2, bookmark="bm_s5_1")
    add_heading_rtl(doc, "5.1.1  التوسعات التقنية", 3)
    add_list(
        doc,
        [
            "طابور استدلال للرؤية عند ارتفاع الحمل.",
            "تخزين سحابي للصور بدل المحلي.",
            "تطبيق جوال أو PWA للإبلاغ الميداني.",
            "مضلعات أحياء واستعلامات مكانية أغنى.",
            "إشعارات فورية عند تغيّر حالة البلاغ.",
        ],
    )
    add_heading_rtl(doc, "5.1.2  التحسينات الوظيفية", 3)
    add_list(
        doc,
        [
            "جدول دقة النموذج لكل صنف على مجموعة اختبار ثابتة.",
            "مراجعة جدول الشدة مع مهندس إنشائي قبل نشر ميداني واسع.",
            "تحليلات زمنية وجغرافية أغنى.",
            "تصدير PDF للبلاغ المعتمد وسجل المراجعة.",
        ],
    )

    add_heading_rtl(doc, "5.2  الصعوبات التي واجهناها", 2, bookmark="bm_s5_2")
    add_heading_rtl(doc, "5.2.1  تحديات تقنية", 3)
    add_list(
        doc,
        [
            "ضبط RBAC على الخادم لا على إخفاء الأزرار فقط.",
            "تشغيل YOLOv8 داخل Node عبر ONNX مع letterbox وفك المخرجات.",
            "تجميع كشف متعدد الصور دون تضخيم الشدة بكشوف ضعيفة.",
            "ضبط CORS بين Vite وNestJS.",
            "توحيد العقود عبر packages/shared.",
        ],
    )
    add_heading_rtl(doc, "5.2.2  تحديات منهجية", 3)
    add_list(
        doc,
        [
            "تحديد النطاق: فرز وتوثيق لا تقرير إنشائي موقّع.",
            "التوازن بين الدراسة والعمل المشترك.",
            "مقاومة توسّع النطاق حول مسار بلاغ → مراجعة → مساعدة.",
            "مزامنة توثيق الأطروحة واللقطات مع تطوّر الواجهة.",
        ],
    )

    add_para(
        doc,
        "خلاصة الفصل: عرضنا مسارات تطوير لاحقة، واستعرضنا بشفافية التحديات التي واجهناها كفريق ثنائي وكيف صقلت تجربتنا.",
        bold=True,
    )

    # ===================== REFERENCES =====================
    chapter_title(doc, "المراجع والمصادر", bookmark="bm_refs")
    refs = [
        "1. NestJS Official Documentation — https://docs.nestjs.com",
        "2. React Documentation — https://react.dev",
        "3. Prisma ORM Documentation — https://www.prisma.io/docs",
        "4. PostgreSQL Official Documentation — https://www.postgresql.org/docs",
        "5. Vite Documentation — https://vitejs.dev",
        "6. Tailwind CSS Documentation — https://tailwindcss.com/docs",
        "7. JWT Introduction — https://jwt.io/introduction",
        "8. TanStack React Query — https://tanstack.com/query",
        "9. Leaflet Documentation — https://leafletjs.com",
        "10. Ultralytics YOLOv8 — https://docs.ultralytics.com",
        "11. ONNX Runtime — https://onnxruntime.ai",
        "12. Mayya et al., PECCD dataset — Mendeley Data, 2025",
        "13. Zod Documentation — https://zod.dev",
        "14. Docker Documentation — https://docs.docker.com",
        "15. React Router Documentation — https://reactrouter.com",
        "16. Jest Documentation — https://jestjs.io",
    ]
    for r in refs:
        add_para(doc, r, size=13, space_after=6, justify=False)

    # ===================== ENGLISH ABSTRACT (LAST PAGE) =====================
    page_break(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Abstract")
    style_run(run, size=18, bold=True, english=True)
    add_bookmark(p, "bm_abstract_en")

    add_para(
        doc,
        "This graduation project presents an integrated web platform for documenting property damage after disasters "
        "or localized incidents. We designed a single channel that connects citizens and residents with a supervisor "
        "and aid organizations. A report carries a geographic location, a description, and photographs. A YOLOv8 object "
        "detector, exported to ONNX and executed inside the NestJS API, proposes an initial damage level. A human "
        "supervisor still approves, rejects, or edits the report. Organizations then attach planned, in-progress, or "
        "delivered aid only to approved reports.",
        size=13,
        english=True,
    )
    add_para(
        doc,
        "We built the system as a monorepo: a React client with Vite, Tailwind CSS, and Leaflet; a NestJS API; a shared "
        "package for contracts and detection aggregation; and PostgreSQL accessed through Prisma. Six PECCD defect classes "
        "drive the vision pipeline. Multi-image findings are reduced by a highest-severity rule with a confidence floor "
        "of 0.35 and an approximate box-coverage ratio. Chapter 4 documents the real UI with nine numbered screenshots "
        "covering the citizen, supervisor, and organization workflows.",
        size=13,
        english=True,
    )
    add_para(
        doc,
        "Keywords: damage documentation, YOLOv8, ONNX Runtime, GIS, aid distribution, role-based access control, "
        "React, NestJS.",
        size=13,
        bold=True,
        english=True,
    )

    doc.save(OUT)
    print("saved", OUT)
    print("figures registered", len(FIGURES), "tables registered", len(TABLES))
    update_fields_in_word(OUT)


if __name__ == "__main__":
    build()
